# -*- coding: utf-8 -*-
# @Time    : 2026
# @Author  : 李涛
# @File    : knowledge_graph_engine.py
# @Desc    : 知识图谱引擎基类（纯Python实现 + JSON持久化）
# Copyright (c) 2026 保留所有权利，开源遵循 MIT 开源协议

"""
知识图谱引擎基类

基于纯Python dict实现有向图，存储概念、实体及关系。
支持：
- 节点/边的增删查
- 关键词搜索
- N跳关联查询
- JSON导出/导入
- Excel导入导出
- 持久化到文件
"""

import json
import os
from collections import defaultdict
from typing import Any, Dict, List, Optional, Set

try:
    from openpyxl import Workbook, load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .base_engine import BaseEngine


class KnowledgeGraphEngine(BaseEngine):
    """知识图谱引擎基类"""

    def __init__(self, data_dir: Optional[str] = None):
        """
        Args:
            data_dir: 数据目录路径，存放 entities.json 和 relations.json
        """
        if data_dir is None:
            data_dir = os.path.dirname(__file__)
        self._data_dir = data_dir
        self._entities_path = os.path.join(data_dir, "entities.json")
        self._relations_path = os.path.join(data_dir, "relations.json")

        # 节点: {node_id: {name, type, desc, properties}}
        self._nodes: Dict[str, Dict[str, Any]] = {}
        # 出边: {from_id: {to_id: {relation, properties}}}
        self._out_edges: Dict[str, Dict[str, Dict[str, Any]]] = defaultdict(dict)
        # 入边: {to_id: {from_id: {relation, properties}}}
        self._in_edges: Dict[str, Dict[str, Dict[str, Any]]] = defaultdict(dict)
        self._load()

    # ------------------------------------------------------------------
    # 节点操作
    # ------------------------------------------------------------------

    def add_node(self, node_id: str, name: str, node_type: str = "", desc: str = "",
                 properties: Optional[Dict[str, Any]] = None) -> bool:
        """添加节点"""
        self._nodes[node_id] = {
            "name": name,
            "type": node_type,
            "desc": desc,
            "properties": properties or {},
        }
        return True

    def remove_node(self, node_id: str) -> bool:
        """删除节点及其相关边"""
        if node_id not in self._nodes:
            return False
        del self._nodes[node_id]
        # 删除相关边
        self._out_edges.pop(node_id, None)
        for to_dict in self._out_edges.values():
            to_dict.pop(node_id, None)
        self._in_edges.pop(node_id, None)
        for from_dict in self._in_edges.values():
            from_dict.pop(node_id, None)
        return True

    def get_node(self, node_id: str) -> Optional[Dict[str, Any]]:
        """获取单个节点"""
        if node_id in self._nodes:
            return dict(self._nodes[node_id], id=node_id)
        return None

    def get_all_nodes(self) -> List[Dict[str, Any]]:
        """获取所有节点"""
        return [dict(node, id=nid) for nid, node in self._nodes.items()]

    def get_all(self) -> List[Dict[str, Any]]:
        """获取所有节点（与基类统一）"""
        return self.get_all_nodes()

    def search(self, keyword: str) -> List[Dict[str, Any]]:
        """关键词搜索节点（支持多关键词拆分匹配）"""
        all_nodes = [dict(node, id=nid) for nid, node in self._nodes.items()]

        def get_search_text(node: Dict[str, Any]):
            name_text = f"{node.get('id', '')} {node.get('name', '')}"
            full_text = f"{name_text} {node.get('desc', '')}"
            return name_text, full_text

        return self._multi_keyword_search(all_nodes, get_search_text, keyword)

    def search_nodes(self, keyword: str) -> List[Dict[str, Any]]:
        """关键词搜索节点（别名）"""
        return self.search(keyword)

    def get_node_count(self) -> int:
        """获取节点总数"""
        return len(self._nodes)

    # ------------------------------------------------------------------
    # 边（关系）操作
    # ------------------------------------------------------------------

    def add_edge(self, from_id: str, to_id: str, relation: str,
                 properties: Optional[Dict[str, Any]] = None) -> bool:
        """添加边（关系）"""
        if from_id not in self._nodes or to_id not in self._nodes:
            return False
        self._out_edges[from_id][to_id] = {"relation": relation, "properties": properties or {}}
        self._in_edges[to_id][from_id] = {"relation": relation, "properties": properties or {}}
        return True

    def remove_edge(self, from_id: str, to_id: str) -> bool:
        """删除边（关系）"""
        if to_id in self._out_edges.get(from_id, {}):
            del self._out_edges[from_id][to_id]
            if not self._out_edges[from_id]:
                self._out_edges.pop(from_id, None)
        if from_id in self._in_edges.get(to_id, {}):
            del self._in_edges[to_id][from_id]
            if not self._in_edges[to_id]:
                self._in_edges.pop(to_id, None)
        return True

    def get_edges(self) -> List[Dict[str, Any]]:
        """获取所有边"""
        edges = []
        for from_id, to_dict in self._out_edges.items():
            for to_id, data in to_dict.items():
                edges.append({
                    "from": from_id,
                    "to": to_id,
                    "relation": data.get("relation", ""),
                    "properties": data.get("properties", {}),
                })
        return edges

    def get_edge_count(self) -> int:
        """获取边总数"""
        return sum(len(to_dict) for to_dict in self._out_edges.values())

    # ------------------------------------------------------------------
    # 关联查询
    # ------------------------------------------------------------------

    def query_related(self, node_id: str, depth: int = 1) -> Dict[str, Any]:
        """N跳关联查询
        
        Args:
            node_id: 中心节点ID
            depth: 查询深度（跳数）
            
        Returns:
            {center: 中心节点信息, relations: 关联关系列表}
        """
        center = self.get_node(node_id)
        if center is None:
            return {"center": None, "relations": []}

        visited: Set[str] = {node_id}
        current_level = {node_id}
        relations = []

        for _ in range(depth):
            next_level: Set[str] = set()
            for nid in current_level:
                # 出边
                for to_id, data in self._out_edges.get(nid, {}).items():
                    if to_id not in visited:
                        visited.add(to_id)
                        next_level.add(to_id)
                    relations.append({
                        "from": nid,
                        "to": to_id,
                        "from_name": self._nodes[nid].get("name", nid),
                        "to_name": self._nodes[to_id].get("name", to_id),
                        "relation": data.get("relation", ""),
                    })
                # 入边
                for from_id, data in self._in_edges.get(nid, {}).items():
                    if from_id not in visited:
                        visited.add(from_id)
                        next_level.add(from_id)
                    relations.append({
                        "from": from_id,
                        "to": nid,
                        "from_name": self._nodes[from_id].get("name", from_id),
                        "to_name": self._nodes[nid].get("name", nid),
                        "relation": data.get("relation", ""),
                    })
            current_level = next_level

        return {"center": center, "relations": relations}

    # ------------------------------------------------------------------
    # 持久化
    # ------------------------------------------------------------------

    def save(self) -> None:
        """保存到JSON文件"""
        nodes = [dict(node, id=nid) for nid, node in self._nodes.items()]
        with open(self._entities_path, "w", encoding="utf-8") as f:
            json.dump({"nodes": nodes}, f, ensure_ascii=False, indent=2)

        edges = self.get_edges()
        with open(self._relations_path, "w", encoding="utf-8") as f:
            json.dump({"relations": edges}, f, ensure_ascii=False, indent=2)

    def _load(self) -> None:
        """从JSON文件加载"""
        if os.path.exists(self._entities_path):
            try:
                with open(self._entities_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                for node in data.get("nodes", []):
                    node_id = node.pop("id")
                    self._nodes[node_id] = node
            except Exception as e:
                print(f"[KnowledgeGraphEngine] 加载节点失败: {e}")

        if os.path.exists(self._relations_path):
            try:
                with open(self._relations_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                for edge in data.get("relations", []):
                    self._out_edges[edge["from"]][edge["to"]] = {
                        "relation": edge.get("relation", ""),
                        "properties": edge.get("properties", {}),
                    }
                    self._in_edges[edge["to"]][edge["from"]] = {
                        "relation": edge.get("relation", ""),
                        "properties": edge.get("properties", {}),
                    }
            except Exception as e:
                print(f"[KnowledgeGraphEngine] 加载边失败: {e}")

    # ------------------------------------------------------------------
    # 导入导出
    # ------------------------------------------------------------------

    def export_to_json(self) -> str:
        """导出为JSON字符串"""
        data = {
            "nodes": [dict(node, id=nid) for nid, node in self._nodes.items()],
            "relations": self.get_edges(),
        }
        return json.dumps(data, ensure_ascii=False, indent=2)

    def import_from_json(self, json_str: str) -> bool:
        """从JSON字符串导入"""
        try:
            data = json.loads(json_str)
            self._nodes.clear()
            self._out_edges.clear()
            self._in_edges.clear()

            for node in data.get("nodes", []):
                node_id = node.pop("id")
                self._nodes[node_id] = node

            for edge in data.get("relations", []):
                self._out_edges[edge["from"]][edge["to"]] = {
                    "relation": edge.get("relation", ""),
                    "properties": edge.get("properties", {}),
                }
                self._in_edges[edge["to"]][edge["from"]] = {
                    "relation": edge.get("relation", ""),
                    "properties": edge.get("properties", {}),
                }
            self.save()
            return True
        except Exception as e:
            print(f"[KnowledgeGraphEngine] 导入失败: {e}")
            return False

    def export_to_excel(self, file_path: str) -> bool:
        """导出知识图谱到Excel文件"""
        if not OPENPYXL_AVAILABLE:
            print("[KnowledgeGraphEngine] 导出失败: 未安装 openpyxl")
            return False

        try:
            wb = Workbook()

            ws_nodes = wb.active
            ws_nodes.title = "实体"
            ws_nodes.append(["id", "name", "type", "desc", "properties"])

            for nid, node in self._nodes.items():
                properties_str = json.dumps(node.get("properties", {}), ensure_ascii=False)
                ws_nodes.append([
                    nid,
                    node.get("name", ""),
                    node.get("type", ""),
                    node.get("desc", ""),
                    properties_str,
                ])

            ws_relations = wb.create_sheet(title="关系")
            ws_relations.append(["from", "to", "relation", "properties"])

            for from_id, to_dict in self._out_edges.items():
                for to_id, data in to_dict.items():
                    properties_str = json.dumps(data.get("properties", {}), ensure_ascii=False)
                    from_name = self._nodes.get(from_id, {}).get("name", from_id)
                    to_name = self._nodes.get(to_id, {}).get("name", to_id)
                    ws_relations.append([
                        from_name,
                        to_name,
                        data.get("relation", ""),
                        properties_str,
                    ])

            wb.save(file_path)
            return True
        except Exception as e:
            print(f"[KnowledgeGraphEngine] 导出Excel失败: {e}")
            return False

    def import_from_excel(self, file_path: str) -> bool:
        """从Excel文件导入知识图谱"""
        if not OPENPYXL_AVAILABLE:
            print("[KnowledgeGraphEngine] 导入失败: 未安装 openpyxl")
            return False

        try:
            wb = load_workbook(file_path, data_only=True)

            self._nodes.clear()
            self._out_edges.clear()
            self._in_edges.clear()

            if "实体" in wb.sheetnames:
                ws_nodes = wb["实体"]
                header = [cell.value for cell in next(ws_nodes.iter_rows(min_row=1, max_row=1))]
                if header != ["id", "name", "type", "desc", "properties"]:
                    print("[KnowledgeGraphEngine] 实体表格式错误")
                    return False

                for row in ws_nodes.iter_rows(min_row=2):
                    row_values = [cell.value if cell.value is not None else "" for cell in row]
                    if len(row_values) >= 2 and row_values[0]:
                        node_id = str(row_values[0])
                        properties = {}
                        if row_values[4]:
                            try:
                                properties = json.loads(row_values[4])
                            except:
                                properties = {}
                        self._nodes[node_id] = {
                            "name": row_values[1],
                            "type": row_values[2],
                            "desc": row_values[3],
                            "properties": properties,
                        }

            if "关系" in wb.sheetnames:
                ws_relations = wb["关系"]
                header = [cell.value for cell in next(ws_relations.iter_rows(min_row=1, max_row=1))]
                if header != ["from", "to", "relation", "properties"]:
                    print("[KnowledgeGraphEngine] 关系表格式错误")
                    return False

                # 建立 name -> id 映射
                name_to_id = {}
                for nid, node in self._nodes.items():
                    name = node.get("name", "")
                    if name:
                        name_to_id[name] = nid

                for row in ws_relations.iter_rows(min_row=2):
                    row_values = [cell.value if cell.value is not None else "" for cell in row]
                    if len(row_values) >= 3 and row_values[0] and row_values[1]:
                        from_name = str(row_values[0])
                        to_name = str(row_values[1])
                        relation = row_values[2]
                        properties = {}
                        if len(row_values) >= 4 and row_values[3]:
                            try:
                                properties = json.loads(row_values[3])
                            except:
                                properties = {}

                        # 根据名称查找id
                        from_id = name_to_id.get(from_name, from_name)
                        to_id = name_to_id.get(to_name, to_name)

                        if from_id in self._nodes and to_id in self._nodes:
                            self._out_edges[from_id][to_id] = {
                                "relation": relation,
                                "properties": properties,
                            }
                            self._in_edges[to_id][from_id] = {
                                "relation": relation,
                                "properties": properties,
                            }
                        else:
                            print(f"[KnowledgeGraphEngine] 关系跳过: 节点不存在 from={from_name}, to={to_name}")

            self.save()
            return True
        except Exception as e:
            print(f"[KnowledgeGraphEngine] 导入Excel失败: {e}")
            return False