# -*- coding: utf-8 -*-
# @Time    : 2026
# @Author  : 李涛
# @File    : case_engine.py
# @Desc    : 案例引擎基类（YAML加载 + 查询/匹配）
# Copyright (c) 2026 保留所有权利，开源遵循 MIT 开源协议

"""
案例引擎基类

管理领域参考案例（工程实例、设计方案、分析报告等）。
支持：
- 案例的增删改查
- 关键词搜索
- 分类筛选
- YAML持久化
- Excel导入导出
"""

import os
import copy
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional

import yaml

try:
    from openpyxl import Workbook, load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .base_engine import BaseEngine


class CaseEngine(BaseEngine):
    """案例引擎基类"""

    def __init__(self, cases_path: Optional[str] = None, default_cases: Optional[List[Dict]] = None):
        """
        Args:
            cases_path: 案例文件路径
            default_cases: 默认案例列表
        """
        if cases_path is None:
            cases_path = os.path.join(os.path.dirname(__file__), "cases.yaml")
        self._cases_path = cases_path
        self._cases: List[Dict[str, Any]] = []
        self._default_cases = default_cases or []
        self._load()

    def add_case(self, case: Dict[str, Any]) -> str:
        """添加案例

        Args:
            case: 案例字典，需包含 title, category, description

        Returns:
            案例ID
        """
        case_id = case.get("id") or self._generate_id()
        new_case = {
            "id": case_id,
            "title": case.get("title", ""),
            "category": case.get("category", ""),
            "description": case.get("description", ""),
            "vehicle_type": case.get("vehicle_type", ""),
            "key_params": case.get("key_params", ""),
            "results": case.get("results", ""),
            "source": case.get("source", ""),
            "created_at": case.get("created_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            "updated_at": case.get("updated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        }
        self._cases.append(new_case)
        return case_id

    def update_case(self, case_id: str, updates: Dict[str, Any]) -> bool:
        """更新案例

        Args:
            case_id: 案例ID
            updates: 要更新的字段

        Returns:
            是否成功更新
        """
        for case in self._cases:
            if case["id"] == case_id:
                for key in ("title", "category", "description", "vehicle_type", "key_params", "results", "source"):
                    if key in updates:
                        case[key] = updates[key]
                case["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                return True
        return False

    def delete_case(self, case_id: str) -> bool:
        """删除案例

        Args:
            case_id: 案例ID

        Returns:
            是否成功删除
        """
        for i, case in enumerate(self._cases):
            if case["id"] == case_id:
                self._cases.pop(i)
                return True
        return False

    def get_case(self, case_id: str) -> Optional[Dict[str, Any]]:
        """获取单个案例"""
        for case in self._cases:
            if case["id"] == case_id:
                return copy.deepcopy(case)
        return None

    def get_all(self) -> List[Dict[str, Any]]:
        """获取所有案例"""
        return copy.deepcopy(self._cases)

    def search(self, keyword: str) -> List[Dict[str, Any]]:
        """关键词搜索案例（支持多关键词拆分匹配title、description、vehicle_type、key_params、results）"""
        all_cases = copy.deepcopy(self._cases)

        def get_search_text(case: Dict[str, Any]):
            name_text = case.get('title', '')
            full_text = f"{name_text} {case.get('description', '')} {case.get('vehicle_type', '')} {case.get('key_params', '')} {case.get('results', '')}"
            return name_text, full_text

        return self._multi_keyword_search(all_cases, get_search_text, keyword)

    def get_by_category(self, category: str) -> List[Dict[str, Any]]:
        """按分类获取案例"""
        if category == "全部" or not category:
            return copy.deepcopy(self._cases)
        return [copy.deepcopy(c) for c in self._cases if c.get("category") == category]

    def list_categories(self) -> List[str]:
        """列出所有分类"""
        categories = set()
        for case in self._cases:
            cat = case.get("category", "")
            if cat:
                categories.add(cat)
        return sorted(categories)

    def get_case_count(self) -> int:
        """获取案例总数"""
        return len(self._cases)

    def save(self) -> None:
        """保存到YAML文件"""
        with open(self._cases_path, "w", encoding="utf-8") as f:
            yaml.dump(
                {"cases": self._cases},
                f,
                allow_unicode=True,
                default_flow_style=False,
                sort_keys=False,
            )

    def export_to_excel(self, file_path: str) -> bool:
        """导出案例到Excel文件"""
        if not OPENPYXL_AVAILABLE:
            print("[CaseEngine] 导出失败: 未安装 openpyxl")
            return False

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "案例"
            ws.append(["id", "title", "category", "description", "vehicle_type", "key_params", "results", "source"])

            for case in self._cases:
                ws.append([
                    case.get("id", ""),
                    case.get("title", ""),
                    case.get("category", ""),
                    case.get("description", ""),
                    case.get("vehicle_type", ""),
                    case.get("key_params", ""),
                    case.get("results", ""),
                    case.get("source", ""),
                ])

            wb.save(file_path)
            return True
        except Exception as e:
            print(f"[CaseEngine] 导出Excel失败: {e}")
            return False

    def import_from_excel(self, file_path: str) -> bool:
        """从Excel文件导入案例"""
        if not OPENPYXL_AVAILABLE:
            print("[CaseEngine] 导入失败: 未安装 openpyxl")
            return False

        try:
            wb = load_workbook(file_path, data_only=True)

            if "案例" not in wb.sheetnames:
                print("[CaseEngine] 案例表不存在")
                return False

            ws = wb["案例"]
            header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            if header != ["id", "title", "category", "description", "vehicle_type", "key_params", "results", "source"]:
                print("[CaseEngine] 案例表格式错误")
                return False

            self._cases.clear()

            for row in ws.iter_rows(min_row=2):
                row_values = [cell.value if cell.value is not None else "" for cell in row]
                if len(row_values) >= 3 and row_values[1]:
                    case_id = str(row_values[0]) if row_values[0] else self._generate_id()
                    self._cases.append({
                        "id": case_id,
                        "title": row_values[1],
                        "category": row_values[2],
                        "description": row_values[3],
                        "vehicle_type": row_values[4],
                        "key_params": row_values[5],
                        "results": row_values[6],
                        "source": row_values[7],
                        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })

            self.save()
            return True
        except Exception as e:
            print(f"[CaseEngine] 导入Excel失败: {e}")
            return False

    def _load(self) -> None:
        """从YAML文件加载"""
        if os.path.exists(self._cases_path):
            try:
                with open(self._cases_path, "r", encoding="utf-8") as f:
                    data = yaml.safe_load(f)
                if data and "cases" in data:
                    self._cases = data["cases"]
            except Exception as e:
                print(f"[CaseEngine] 加载案例失败: {e}")

        if not self._cases and self._default_cases:
            self._cases = copy.deepcopy(self._default_cases)
            for case in self._cases:
                case["created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                case["updated_at"] = case["created_at"]
            self.save()

    @staticmethod
    def _generate_id() -> str:
        """生成唯一案例ID"""
        return f"case_{uuid.uuid4().hex[:8]}"