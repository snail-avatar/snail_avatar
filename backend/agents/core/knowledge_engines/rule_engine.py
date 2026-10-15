# -*- coding: utf-8 -*-
# @Time    : 2026
# @Author  : 李涛
# @File    : rule_engine.py
# @Desc    : 规则引擎基类（YAML加载 + 查询/匹配）
# Copyright (c) 2026 保留所有权利，开源遵循 MIT 开源协议

"""
规则引擎基类

管理领域规则（物理公式、经验公式、设计准则）。
支持：
- 规则的增删改查
- 关键词搜索
- 分类筛选
- YAML持久化
- Excel导入导出
"""

import os
import copy
import uuid
from datetime import datetime
from typing import Any, Dict, List, Optional

import yaml

try:
    from openpyxl import Workbook, load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .base_engine import BaseEngine


class RuleEngine(BaseEngine):
    """规则引擎基类"""

    def __init__(self, rules_path: Optional[str] = None, default_rules: Optional[List[Dict]] = None):
        """
        Args:
            rules_path: 规则文件路径
            default_rules: 默认规则列表
        """
        if rules_path is None:
            rules_path = os.path.join(os.path.dirname(__file__), "rules.yaml")
        self._rules_path = rules_path
        self._rules: List[Dict[str, Any]] = []
        self._default_rules = default_rules or []
        self._load()

    def add_rule(self, rule: Dict[str, Any]) -> str:
        """添加规则

        Args:
            rule: 规则字典，需包含 name, category, description

        Returns:
            规则ID
        """
        rule_id = rule.get("id") or self._generate_id()
        new_rule = {
            "id": rule_id,
            "name": rule.get("name", ""),
            "category": rule.get("category", ""),
            "description": rule.get("description", ""),
            "formula": rule.get("formula", ""),
            "source": rule.get("source", ""),
            "created_at": rule.get("created_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            "updated_at": rule.get("updated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        }
        self._rules.append(new_rule)
        return rule_id

    def update_rule(self, rule_id: str, updates: Dict[str, Any]) -> bool:
        """更新规则

        Args:
            rule_id: 规则ID
            updates: 要更新的字段

        Returns:
            是否成功更新
        """
        for rule in self._rules:
            if rule["id"] == rule_id:
                for key in ("name", "category", "description", "formula", "source"):
                    if key in updates:
                        rule[key] = updates[key]
                rule["updated_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                return True
        return False

    def delete_rule(self, rule_id: str) -> bool:
        """删除规则

        Args:
            rule_id: 规则ID

        Returns:
            是否成功删除
        """
        for i, rule in enumerate(self._rules):
            if rule["id"] == rule_id:
                self._rules.pop(i)
                return True
        return False

    def get_rule(self, rule_id: str) -> Optional[Dict[str, Any]]:
        """获取单个规则"""
        for rule in self._rules:
            if rule["id"] == rule_id:
                return copy.deepcopy(rule)
        return None

    def get_all(self) -> List[Dict[str, Any]]:
        """获取所有规则"""
        return copy.deepcopy(self._rules)

    def search(self, keyword: str) -> List[Dict[str, Any]]:
        """关键词搜索规则（支持多关键词拆分匹配name、description、formula）"""
        all_rules = copy.deepcopy(self._rules)

        def get_search_text(rule: Dict[str, Any]):
            name_text = rule.get('name', '')
            full_text = f"{name_text} {rule.get('description', '')} {rule.get('formula', '')}"
            return name_text, full_text

        return self._multi_keyword_search(all_rules, get_search_text, keyword)

    def get_by_category(self, category: str) -> List[Dict[str, Any]]:
        """按分类获取规则"""
        if category == "全部" or not category:
            return copy.deepcopy(self._rules)
        return [copy.deepcopy(r) for r in self._rules if r.get("category") == category]

    def list_categories(self) -> List[str]:
        """列出所有分类"""
        categories = set()
        for rule in self._rules:
            cat = rule.get("category", "")
            if cat:
                categories.add(cat)
        return sorted(categories)

    def get_rule_count(self) -> int:
        """获取规则总数"""
        return len(self._rules)

    def save(self) -> None:
        """保存到YAML文件"""
        with open(self._rules_path, "w", encoding="utf-8") as f:
            yaml.dump(
                {"rules": self._rules},
                f,
                allow_unicode=True,
                default_flow_style=False,
                sort_keys=False,
            )

    def export_to_excel(self, file_path: str) -> bool:
        """导出规则到Excel文件"""
        if not OPENPYXL_AVAILABLE:
            print("[RuleEngine] 导出失败: 未安装 openpyxl")
            return False

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "规则"
            ws.append(["id", "name", "category", "description", "formula", "source"])

            for rule in self._rules:
                ws.append([
                    rule.get("id", ""),
                    rule.get("name", ""),
                    rule.get("category", ""),
                    rule.get("description", ""),
                    rule.get("formula", ""),
                    rule.get("source", ""),
                ])

            wb.save(file_path)
            return True
        except Exception as e:
            print(f"[RuleEngine] 导出Excel失败: {e}")
            return False

    def import_from_excel(self, file_path: str) -> bool:
        """从Excel文件导入规则"""
        if not OPENPYXL_AVAILABLE:
            print("[RuleEngine] 导入失败: 未安装 openpyxl")
            return False

        try:
            wb = load_workbook(file_path, data_only=True)

            if "规则" not in wb.sheetnames:
                print("[RuleEngine] 规则表不存在")
                return False

            ws = wb["规则"]
            header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            if header != ["id", "name", "category", "description", "formula", "source"]:
                print("[RuleEngine] 规则表格式错误")
                return False

            self._rules.clear()

            for row in ws.iter_rows(min_row=2):
                row_values = [cell.value if cell.value is not None else "" for cell in row]
                if len(row_values) >= 3 and row_values[1]:
                    rule_id = str(row_values[0]) if row_values[0] else self._generate_id()
                    self._rules.append({
                        "id": rule_id,
                        "name": row_values[1],
                        "category": row_values[2],
                        "description": row_values[3],
                        "formula": row_values[4],
                        "source": row_values[5],
                        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    })

            self.save()
            return True
        except Exception as e:
            print(f"[RuleEngine] 导入Excel失败: {e}")
            return False

    def _load(self) -> None:
        """从YAML文件加载"""
        if os.path.exists(self._rules_path):
            try:
                with open(self._rules_path, "r", encoding="utf-8") as f:
                    data = yaml.safe_load(f)
                if data and "rules" in data:
                    self._rules = data["rules"]
            except Exception as e:
                print(f"[RuleEngine] 加载规则失败: {e}")

        if not self._rules and self._default_rules:
            self._rules = copy.deepcopy(self._default_rules)
            for rule in self._rules:
                rule["created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                rule["updated_at"] = rule["created_at"]
            self.save()

    @staticmethod
    def _generate_id() -> str:
        """生成唯一规则ID"""
        return f"rule_{uuid.uuid4().hex[:8]}"